"""
Разбор сменного отчёта: дубль, непрочитанное, итоги месяца, сверка.

Четыре вещи, которые разбор обязан делать вслух:

  1. одна и та же смена на двух листах — помечается, в суммы идёт
     одна, но из базы не исчезает ни одна;
  2. простой, у которого время в ячейке есть, а минут из него не
     вышло, — ложится в «не разобрано» с листом, строкой и причиной,
     а не превращается в ноль;
  3. итоги месяца (хвост листа ниже «Средних значений») переносятся
     как есть, включая ноль: «выполнение плана 0 %» значит ровно то,
     что написано;
  4. отчёт сверки показывает всё это ДО сохранения.

У каждой проверки «не теряет» есть парная «не удваивает» или «не
шумит»: разбор, который считает дубль дважды, и разбор, который
жалуется на ровном файле, одинаково бесполезны — первому не верят в
сумме, второму не верят вообще.

Запуск: python tests/test_xls_reconcile.py
"""

from pathlib import Path
import tempfile

from sandbox import Sandbox, check, finish

sb = Sandbox()

import openpyxl

from backend.services.production_report_import import read_workbook
from backend.services.production_import_service import (
    get_connection, reconciliation, save_workbook,
)

YEAR = 2033
folder = Path(tempfile.mkdtemp(prefix="xls_rec_"))

HEAD = [
    ("Смена", ""), ("№", ""), ("Начальник смены", ""),
    ("Участок формовки", "план"), ("Участок формовки", "факт"),
    ("Простой оборудования формовка", "мин"),
    ("Простой оборудования формовка", "причина и замечание"),
    ("Простой оборудования массоподготовка", "мин"),
    ("Простой оборудования массоподготовка", "причина и замечание"),
    ("Простой оборудования печь и сушилка", "мин"),
    ("Простой оборудования печь и сушилка", "причина и замечание"),
    ("Простой оборудования упаковка", "мин"),
    ("Простой оборудования упаковка", "причина и замечание"),
    ("1,4 НФ полнател", "штук"),
]


def sheet_for(book, title, rows, tail=None, first=None):
    sheet = book.create_sheet(title)
    for index, (top, sub) in enumerate(HEAD, start=1):
        sheet.cell(1, index, top)
        if sub:
            sheet.cell(2, index, sub)
    line = 3
    if first:                       # первая строка — ночь прошлого месяца
        sheet.cell(line, 1, first[0])
        sheet.cell(line, 3, "Петров")
        sheet.cell(line, 5, first[1])
        sheet.cell(line, 14, first[1])
        line += 1
    total = sum(value for _, value in rows) + (first[1] if first else 0)
    for label, value in rows:
        sheet.cell(line, 1, label)
        sheet.cell(line, 3, "Иванов")
        sheet.cell(line, 5, value)
        sheet.cell(line, 14, value)
        line += 1
    sheet.cell(line, 1, "Средние значения")
    sheet.cell(line, 14, total)
    line += 1
    for label, value, unit in tail or []:
        sheet.cell(line, 25, label)
        if value is not None:
            sheet.cell(line, 26, value)
        if unit:
            sheet.cell(line, 27, unit)
        line += 1
    return sheet


def build(path):
    book = openpyxl.Workbook()
    book.remove(book.active)
    # Июнь: 30 ночь — «своя» и заполненная.
    sheet_for(book, "Июнь", [("1 день", 100), ("30 день", 200), ("30 ночь", 300)],
              tail=[("выполнение плана", 0, "%"), ("Остаток сырца 1,4 НФ в 1С-", None, None)])
    # Июль: первой строкой та же «30 ночь» июня — дубль.
    sheet_for(book, "Июль", [("1 день", 400)], first=("30 ночь", 300))
    book.save(path)
    return path


data = read_workbook(str(build(folder / "dup.xlsx")), YEAR)
report = reconciliation(data)

# ─────────────────────────────────────────────────────────
print("\n1. Дубль смены: одна в суммах, обе в базе")

check("дубль найден", len(report["duplicates"]) == 1, report["duplicates"])

# Проверка не должна ронять весь файл: сбой в одном месте не повод
# скрыть остальные — иначе, чиня по одному, не увидишь всей картины.
first_dup = report["duplicates"][0] if report["duplicates"] else {}
check("это именно 30 июня, ночь",
      (first_dup.get("date") or "").endswith("-06-30")
      and first_dup.get("shift") == "night", report["duplicates"])
check("в суммы идёт на одну меньше",
      report["counted"] == report["shifts"] - 1,
      (report["counted"], report["shifts"]))
check("сказано, какой лист главный",
      "Июнь" in (first_dup.get("conflict") or ""), first_dup.get("conflict"))

june = report["by_month"].get(f"{YEAR}-06") or {}
check("июньская сумма без дубля", june.get("forming") == 600.0, june)

save_workbook(data, "dup.xlsx", "Проверка")
conn = get_connection()
kept = conn.execute(
    "SELECT COUNT(*) FROM xls_shifts WHERE year = ? AND date = ? AND shift = 'night'",
    (YEAR, f"{YEAR}-06-30")).fetchone()[0]
marked = conn.execute(
    "SELECT COUNT(*) FROM xls_shifts WHERE year = ? AND dup_of IS NOT NULL",
    (YEAR,)).fetchone()[0]
conn.close()

# Парная сторона: обе записи в базе, но помечена ровно одна. Если бы
# дубль просто выбрасывали, первая проверка была бы зелёной, а строка
# из файла — потеряна.
check("в базе обе записи", kept == 2, kept)
check("помечена ровно одна", marked == 1, marked)

# ─────────────────────────────────────────────────────────
print("\n1б. Пустая строка проигрывает заполненной")

# В живом файле июньская строка 30.06 пустая, а смену занесли строкой
# «Июля». Правило «главная — со своего листа» в одиночку выбросило бы
# из июньской суммы восемь вагонеток: на боевом июнь даёт 448, а такой
# разбор давал 440. Проверка именно про это.
empty = openpyxl.Workbook()
empty.remove(empty.active)
june = empty.create_sheet("Июнь")
for index, (top, sub) in enumerate(HEAD, start=1):
    june.cell(1, index, top)
    if sub:
        june.cell(2, index, sub)
june.cell(3, 1, "1 день");  june.cell(3, 3, "Иванов")
june.cell(3, 5, 100);       june.cell(3, 14, 100)
june.cell(4, 1, "30 день"); june.cell(4, 3, "Иванов")
june.cell(4, 5, 200);       june.cell(4, 14, 200)
june.cell(5, 1, "30 ночь")                    # строка есть, чисел нет
june.cell(5, 2, "30Н")
# В живом файле в такой строке выпуск записан текстом — «114.704».
# Числом он не становится, но строку «непустой» делает.
june.cell(5, 14, "114.704")
june.cell(6, 1, "Средние значения"); june.cell(6, 14, 300)
sheet_for(empty, "Июль", [("1 день", 400)], first=("30 ночь", 300))
empty.save(folder / "empty_twin.xlsx")

twin = read_workbook(str(folder / "empty_twin.xlsx"), YEAR)
pair = [item for month in twin["months"] for item in month["shifts"]
        if item["date"].endswith("-06-30") and item["shift"] == "night"]
kept_row = [item for item in pair if not item.get("dup_of")]
check("обе записи на месте", len(pair) == 2, pair)
check("в суммы идёт заполненная",
      len(kept_row) == 1 and kept_row[0].get("forming_fact") == 300, kept_row)

twin_report = reconciliation(twin)
june_box = twin_report["by_month"].get(f"{YEAR}-06") or {}
check("июньская сумма не потеряла смену",
      june_box.get("forming") == 600.0, june_box)

# Парная сторона: когда заполнены обе, решает «свой» лист — иначе
# правило «берём заполненную» стало бы «берём последнюю попавшуюся».
both = [item for item in pair if item.get("dup_of")]
check("незаполненная помечена повтором",
      len(both) == 1 and both[0].get("forming_fact") is None, both)
check("а при двух заполненных главной остаётся своя",
      (report["duplicates"][0]["sheet"] or "") != "Июнь"
      if report["duplicates"] else False,
      report["duplicates"])

# ─────────────────────────────────────────────────────────
print("\n2. Простой без минут — в «не разобрано», а не в ноль")

book = openpyxl.load_workbook(folder / "dup.xlsx")
sheet = book["Июнь"]
sheet.cell(4, 6, "04:30:00")             # одиночное время
sheet.cell(4, 7, "Замена ленты")
sheet.cell(5, 6, "10:00-10:40")          # обычный интервал
sheet.cell(5, 7, "Настройка")
book.save(folder / "downtime.xlsx")

data2 = read_workbook(str(folder / "downtime.xlsx"), YEAR)
report2 = reconciliation(data2)
check("простой с непонятным временем посчитан отдельно",
      report2["downtime_unparsed"] == 1, report2["downtime_unparsed"])

rows = [item for month in data2["months"] for item in month["downtime"]]
lone = [item for item in rows if item["minutes"] is None]
check("минуты не выдуманы", lone and lone[0]["interval"] == "04:30:00", lone)
check("причина сохранена", lone and lone[0]["reason"] == "Замена ленты", lone)

# Парная сторона: понятный интервал по-прежнему считается.
good = [item for item in rows if item["minutes"] == 40]
check("обычный интервал считается", len(good) == 1, [x["interval"] for x in rows])

save_workbook(data2, "downtime.xlsx", "Проверка")
conn = get_connection()
unparsed = [dict(row) for row in conn.execute(
    "SELECT sheet, row, raw, reason FROM xls_unparsed WHERE year = ? AND raw IS NOT NULL",
    (YEAR,))]
conn.close()
check("непрочитанное записано с листом и строкой",
      any(item["raw"] == "04:30:00" and item["sheet"] and item["row"]
          for item in unparsed), unparsed[:3])
check("и с причиной, а не просто числом",
      any("не разобрано" in (item["reason"] or "") for item in unparsed), unparsed[:3])

# ─────────────────────────────────────────────────────────
print("\n3. Итоги месяца: переносятся, включая ноль")

totals = data["month_totals"]
check("итоги месяца прочитаны", totals, totals)
check("ноль не потерялся",
      any(item.get("value_num") == 0.0 and "выполнение" in item["title"]
          for item in totals),
      [(item["title"], item.get("value_num")) for item in totals])
check("единица сохранена",
      any(item.get("unit") == "%" for item in totals),
      [(item["title"], item.get("unit")) for item in totals])
check("подпись без числа тоже сохранена",
      any("Остаток сырца" in item["title"] for item in totals),
      [item["title"] for item in totals])

conn = get_connection()
stored = conn.execute(
    "SELECT COUNT(*) FROM xls_month_totals WHERE year = ?", (YEAR,)).fetchone()[0]
conn.close()
check("и записаны в базу", stored > 0, stored)

# ─────────────────────────────────────────────────────────
print("\n4. Отчёт сверки")

for field in ("sheets", "shifts", "counted", "downtime", "downtime_unparsed",
              "notes", "new_values", "month_totals", "by_month",
              "mismatches", "unknown_labels", "new_columns"):
    check(f"в отчёте есть «{field}»", field in report2, sorted(report2))

check("суммы разложены по месяцам",
      all(len(key) == 7 for key in report2["by_month"]), report2["by_month"])

# Парная сторона: на ровном файле отчёт молчит. Иначе «расхождение»
# станет фоновым шумом и его перестанут читать.
clean_book = openpyxl.Workbook()
clean_book.remove(clean_book.active)
sheet_for(clean_book, "Июнь", [("1 день", 10), ("1 ночь", 20)])
clean_book.save(folder / "clean.xlsx")
clean = reconciliation(read_workbook(str(folder / "clean.xlsx"), YEAR))
check("на ровном файле расхождений нет", not clean["mismatches"], clean["mismatches"])
check("и дублей нет", not clean["duplicates"], clean["duplicates"])
check("и непрочитанных простоев нет",
      clean["downtime_unparsed"] == 0, clean["downtime_unparsed"])

finish("Разбор: дубли, непрочитанное, итоги, сверка")
