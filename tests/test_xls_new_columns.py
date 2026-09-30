"""
Новое в Экселе не пропадает молча.

Файл ведут люди: однажды вставят графу, впишут новый вид изделия,
поставят метку, которой раньше не было. Разбор обязан сказать об этом
вслух, а не тихо пройти мимо — «не показать» хуже «не разобрать»
только на словах; на деле молча потерянное число хуже обоих.

Два случая, заданные владельцем 30.09.2026:

  1. НОВЫЙ ЗАГОЛОВОК колонки. Значения ложатся как есть, с заголовком
     из файла, помечаются «новое, смысл не определён» и попадают в
     список на привязку — а не исчезают при разборе.

  2. НОВОЕ ЗНАЧЕНИЕ в знакомой колонке, например шестой вид изделия
     там, где было пять. Оно не теряется, а если сумма разобранного
     не сходится с итогом листа — расхождение показывается.

Проверяются обе стороны. Парная сторона здесь — не «а вдруг лишнее»,
а «а вдруг оно только в журнале»: значение должно и сохраниться, и
быть помеченным. Тест, который проверяет лишь запись в журнал,
зелёный у разбора, который число выбросил, но про это написал.

Запуск: python tests/test_xls_new_columns.py
"""

from pathlib import Path
import tempfile

from sandbox import Sandbox, check, finish

sb = Sandbox()

import openpyxl

from backend.services.production_report_import import read_workbook
from backend.services.production_import_service import get_connection, save_workbook

YEAR = 2032          # года нет в боевой копии: чужое не трогаем

# Лист по образцу живого: две строки заголовков, строки смен, внизу
# «Средние значения» — суммы по колонкам, с которыми и сверяемся.
BASE = [
    ("Смена", ""), ("№", ""), ("Начальник смены", ""),
    ("Участок формовки", "план"), ("Участок формовки", "факт"),
    ("Простой оборудования формовка", "мин"),
    ("Простой оборудования формовка", "причина и замечание"),
    ("Простой оборудования массоподготовка", "мин"),
    ("Простой оборудования массоподготовка", "причина и замечание"),
    ("Простой оборудования печь", "мин"),
    ("Простой оборудования печь", "причина и замечание"),
    ("Простой оборудования упаковка", "мин"),
    ("Простой оборудования упаковка", "причина и замечание"),
    ("1,4 НФ полнател", "штук"),
]


def build(path, extra_header=None, extra_product=None, lose_row=False):
    """
    Лист-месяц. `extra_header` — новая графа, `extra_product` — новый
    вид изделия, `lose_row` — строка с данными и меткой, которую
    разбор не узнает.
    """
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Сентябрь"

    columns = list(BASE)
    if extra_header:
        columns.append((extra_header, ""))
    if extra_product:
        columns.append((extra_product, "штук"))

    for index, (top, sub) in enumerate(columns, start=1):
        sheet.cell(1, index, top)
        if sub:
            sheet.cell(2, index, sub)

    rows = [("1 день", 100), ("1 ночь", 200), ("2 день", 300)]
    line = 3
    totals = {}
    for label, base in rows:
        sheet.cell(line, 1, label)
        sheet.cell(line, 3, "Иванов")
        sheet.cell(line, 4, 10)
        sheet.cell(line, 5, 9)
        sheet.cell(line, 14, base)
        totals[14] = totals.get(14, 0) + base
        if extra_header:
            sheet.cell(line, 15, base + 7)
        if extra_product:
            col = 16 if extra_header else 15
            sheet.cell(line, col, base * 2)
            totals[col] = totals.get(col, 0) + base * 2
        line += 1

    if lose_row:
        # Метка не «день/ночь», а данные есть — как «№53,» в живом файле.
        sheet.cell(line, 1, "№53,")
        sheet.cell(line, 14, 555)
        totals[14] = totals.get(14, 0) + 555
        line += 1

    sheet.cell(line, 1, "Средние значения")
    for col, total in totals.items():
        sheet.cell(line, col, total)

    book.save(path)
    return path


folder = Path(tempfile.mkdtemp(prefix="xls_new_"))


def values_in_db(year=YEAR):
    conn = get_connection()
    try:
        return [dict(row) for row in conn.execute(
            "SELECT key, title, value_num, value_text, sheet, row "
            "FROM xls_values WHERE year = ?", (year,))]
    finally:
        conn.close()


# ─────────────────────────────────────────────────────────
print("\n1. Новая графа: сохраняется и видна как новая")

path = build(folder / "new_column.xlsx", extra_header="Стреч плёнка мессерси")
data = read_workbook(str(path), YEAR)

said = [item["what"] for item in data["problems"] if "новая графа" in item["what"]]
check("о новой графе сказано вслух",
      any("Стреч плёнка мессерси" in text for text in said), said)
check("и сказано, что смысл не определён",
      any("смысл не определён" in text for text in said), said)

save_workbook(data, "new_column.xlsx", "Проверка")
stored = [item for item in values_in_db() if item["title"] == "Стреч плёнка мессерси"]

# Парная сторона: мало написать в журнал — значения должны лечь в базу.
check("значения новой графы сохранены", len(stored) == 3, len(stored))
check("сохранены именно числа из файла",
      sorted(item["value_num"] for item in stored) == [107.0, 207.0, 307.0],
      sorted(item["value_num"] for item in stored))
check("заголовок сохранён как в файле",
      all(item["title"] == "Стреч плёнка мессерси" for item in stored))
check("помечено новым", all(item["key"] == "new" for item in stored),
      {item["key"] for item in stored})
check("видно, откуда взято",
      all(item["sheet"] == "Сентябрь" and item["row"] for item in stored), stored[:1])

# ─────────────────────────────────────────────────────────
print("\n2. Новый вид изделия: не теряется")

path = build(folder / "new_product.xlsx", extra_product="2,1 НФ клинкер")
data = read_workbook(str(path), YEAR)
products = [item for month in data["months"] for shift in month["shifts"]
            for item in shift.get("products") or []]
titles = {item["title"] for item in products}

check("новый вид попал в выпуск", "2,1 НФ клинкер" in titles, titles)
check("и со своим количеством",
      sorted(item["value"] for item in products if item["title"] == "2,1 НФ клинкер")
      == [200.0, 400.0, 600.0])

# ─────────────────────────────────────────────────────────
print("\n3. Сумма не сошлась с итогом листа — расхождение видно")

path = build(folder / "lost_row.xlsx", lose_row=True)
data = read_workbook(str(path), YEAR)
texts = [item["what"] for item in data["problems"]]

check("о неразобранной метке сказано",
      any("№53," in text and "не разобрал" in text for text in texts), texts)
check("расхождение с итогом листа показано",
      any("расхождение" in text and "555" in text for text in texts), texts)
check("названо, по какой колонке",
      any("1 4 НФ полнател" in text and "расхождение" in text for text in texts), texts)

# Парная сторона: когда всё сошлось, жалоб быть не должно — иначе
# «расхождение» станет фоновым шумом, и его перестанут читать.
clean = read_workbook(str(build(folder / "clean.xlsx")), YEAR)
check("на ровном файле расхождений нет",
      not [item for item in clean["problems"] if "расхождение" in item["what"]],
      [item["what"] for item in clean["problems"]])
check("и новых граф тоже нет",
      not [item for item in clean["problems"] if "новая графа" in item["what"]],
      [item["what"] for item in clean["problems"]])

# ─────────────────────────────────────────────────────────
print("\n4. Брак: все колонки, а не первая")

book = openpyxl.load_workbook(build(folder / "defects.xlsx"))
sheet = book.active
sheet.cell(1, 15, "Брак и отстрел 1,4 НФ полнател")
sheet.cell(2, 15, "штук")
sheet.cell(1, 16, "Брак и отстрел")
sheet.cell(2, 16, "штук")
for line in (3, 4, 5):
    sheet.cell(line, 15, 10)
    sheet.cell(line, 16, 5)
book.save(folder / "defects.xlsx")

data = read_workbook(str(folder / "defects.xlsx"), YEAR)
defects = [shift.get("defect_pieces") for month in data["months"]
           for shift in month["shifts"]]
check("сложены обе колонки брака", defects == [15.0, 15.0, 15.0], defects)

save_workbook(data, "defects.xlsx", "Проверка")
kept = [item for item in values_in_db() if item["key"] == "defect"]
check("каждая колонка брака сохранена отдельно", len(kept) == 6, len(kept))
check("с её собственным заголовком",
      {item["title"] for item in kept} ==
      {"Брак и отстрел 1,4 НФ полнател / штук", "Брак и отстрел / штук"},
      {item["title"] for item in kept})

finish("Новое в Экселе")
